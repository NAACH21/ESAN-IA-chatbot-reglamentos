import os
import re
import sys
import shutil
from typing import List, Dict, Any
from pypdf import PdfReader
from openai import OpenAI

# Compatibilidad SQLite para despliegues en la nube (Linux / Streamlit Cloud)
try:
    __import__("pysqlite3")
    sys.modules["sqlite3"] = sys.modules.pop("pysqlite3")
except ImportError:
    pass

import chromadb

CHROMA_DATA_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "chroma_db")
DOCS_FOLDER = os.path.join(os.path.dirname(os.path.abspath(__file__)), "DOCS")
ARCHIVE_FOLDER = os.path.join(DOCS_FOLDER, "_archivados")
SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".xlsx")

def is_supported_file(filename: str) -> bool:
    lower = filename.lower()
    return lower.endswith(SUPPORTED_EXTENSIONS) and not lower.startswith("~$")

def get_all_supported_files(docs_dir: str = DOCS_FOLDER) -> List[str]:
    """Retorna las rutas completas de todos los archivos soportados recorriendo subcarpetas recursivamente y omitiendo _archivados."""
    found = []
    if not os.path.exists(docs_dir):
        return []
    for root, dirs, files in os.walk(docs_dir):
        # Excluir carpetas temporales, de archivo o del sistema
        dirs[:] = [d for d in dirs if not d.startswith("_") and not d.startswith(".")]
        for file in files:
            if is_supported_file(file):
                found.append(os.path.join(root, file))
    return sorted(found)

def detect_section_or_article(text: str, default_hint: str = "") -> str:
    """Detecta referencias a Títulos, Capítulos, Artículos o Hojas dentro del fragmento."""
    matches = []
    if default_hint:
        matches.append(default_hint)
        
    cap_match = re.search(r'(CAP[ÍI]TULO\s+[IVXLCDM\d]+|T[ÍI]TULO\s+[IVXLCDM\d]+)', text, re.IGNORECASE)
    if cap_match:
        matches.append(cap_match.group(1).upper())
        
    art_matches = re.findall(r'(Art(?:[íi]culo|\.)\s*\d+(?:\s*°|\º)?)', text, re.IGNORECASE)
    if art_matches:
        unique_arts = []
        for a in art_matches:
            clean_a = re.sub(r'\s+', ' ', a).strip()
            if clean_a not in unique_arts:
                unique_arts.append(clean_a)
        matches.append(", ".join(unique_arts[:3]))
        
    if matches:
        return " | ".join(matches)
    return "Sección General / Normativa"

def extract_text_from_pdf(pdf_path: str) -> List[Dict[str, Any]]:
    """Extrae texto página por página de un PDF."""
    reader = PdfReader(pdf_path)
    filename = os.path.basename(pdf_path)
    pages_data = []
    
    for page_num, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        text = text.strip()
        if text:
            text = re.sub(r'\n\s*\n', '\n\n', text)
            pages_data.append({
                "filename": filename,
                "page": page_num,
                "text": text,
                "section_hint": ""
            })
    return pages_data

def extract_text_from_docx(docx_path: str) -> List[Dict[str, Any]]:
    """Extrae párrafos y tablas de un archivo Word (.docx)."""
    import docx
    doc = docx.Document(docx_path)
    filename = os.path.basename(docx_path)
    blocks = []
    current_lines = []
    current_len = 0
    block_num = 1
    current_heading = "Documento Word"

    for para in doc.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        if para.style and "Heading" in para.style.name:
            current_heading = t[:60]
        current_lines.append(t)
        current_len += len(t)
        if current_len >= 1800:
            blocks.append({
                "filename": filename,
                "page": block_num,
                "text": "\n".join(current_lines),
                "section_hint": current_heading
            })
            block_num += 1
            current_lines = []
            current_len = 0

    # Extraer tablas del Word si existen
    for t_idx, table in enumerate(doc.tables, start=1):
        table_lines = []
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                table_lines.append(" | ".join(cells))
        if table_lines:
            current_lines.append(f"[Tabla #{t_idx}]\n" + "\n".join(table_lines))

    if current_lines:
        blocks.append({
            "filename": filename,
            "page": block_num,
            "text": "\n".join(current_lines),
            "section_hint": current_heading
        })
    return blocks

def extract_text_from_xlsx(xlsx_path: str) -> List[Dict[str, Any]]:
    """Extrae hojas y filas estructuradas de un archivo Excel (.xlsx)."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    filename = os.path.basename(xlsx_path)
    sheets_data = []

    for sheet_idx, sheet_name in enumerate(wb.sheetnames, start=1):
        sheet = wb[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        headers = []
        row_texts = []
        for r in rows:
            clean_cells = [str(c).strip() if c is not None else "" for c in r]
            if not any(clean_cells):
                continue
            if not headers:
                headers = [h if h else f"Col_{i+1}" for i, h in enumerate(clean_cells)]
                row_texts.append("Columnas: " + " | ".join(headers))
            else:
                pairs = [f"{headers[i]}: {val}" for i, val in enumerate(clean_cells) if i < len(headers) and val]
                if pairs:
                    row_texts.append(" ; ".join(pairs))

        if row_texts:
            full_sheet_text = f"Hoja de Excel '{sheet_name}':\n" + "\n".join(row_texts)
            sheets_data.append({
                "filename": filename,
                "page": sheet_idx,
                "text": full_sheet_text,
                "section_hint": f"Hoja: {sheet_name}"
            })
    return sheets_data

def extract_document_data(file_path: str) -> List[Dict[str, Any]]:
    """Enruta la extracción según la extensión del archivo (.pdf, .docx, .xlsx)."""
    lower = file_path.lower()
    if lower.endswith(".pdf"):
        return extract_text_from_pdf(file_path)
    elif lower.endswith(".docx"):
        return extract_text_from_docx(file_path)
    elif lower.endswith(".xlsx"):
        return extract_text_from_xlsx(file_path)
    return []

def split_into_chunks(text: str, chunk_size: int = 900, overlap: int = 150) -> List[str]:
    """Divide texto en chunks semánticos con solapamiento."""
    if len(text) <= chunk_size:
        return [text]
    
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        if end >= len(text):
            chunks.append(text[start:].strip())
            break
        
        split_at = text.rfind("\n\n", start, end)
        if split_at == -1 or split_at <= start:
            split_at = text.rfind(". ", start, end)
        if split_at == -1 or split_at <= start:
            split_at = text.rfind("\n", start, end)
        if split_at == -1 or split_at <= start:
            split_at = end
        else:
            split_at += 1
            
        chunk = text[start:split_at].strip()
        if chunk:
            chunks.append(chunk)
        start = max(start + 1, split_at - overlap)
        
    return chunks

def remove_document_from_store(filename: str, docs_dir: str = DOCS_FOLDER) -> int:
    """Elimina los vectores de un documento en ChromaDB y retira el archivo de DOCS."""
    chroma_client = chromadb.PersistentClient(path=CHROMA_DATA_PATH)
    removed_chunks = 0
    try:
        collection = chroma_client.get_collection(name="esan_reglamentos")
        existing = collection.get(where={"source": filename})
        if existing and existing.get("ids"):
            removed_chunks = len(existing["ids"])
            collection.delete(where={"source": filename})
    except Exception:
        pass

    src_path = os.path.join(docs_dir, filename)
    if not os.path.exists(src_path):
        for root, _, files in os.walk(docs_dir):
            if filename in files:
                src_path = os.path.join(root, filename)
                break

    if os.path.exists(src_path):
        os.makedirs(ARCHIVE_FOLDER, exist_ok=True)
        dst_path = os.path.join(ARCHIVE_FOLDER, os.path.basename(filename))
        shutil.move(src_path, dst_path)

    return removed_chunks

def index_single_pdf(api_key: str, pdf_path: str) -> int:
    """Indexa un único archivo nuevo (.pdf, .docx o .xlsx) directamente en ChromaDB."""
    client_ai = OpenAI(api_key=api_key)
    chroma_client = chromadb.PersistentClient(path=CHROMA_DATA_PATH)
    collection = chroma_client.get_or_create_collection(
        name="esan_reglamentos",
        metadata={"hnsw:space": "cosine"}
    )
    filename = os.path.basename(pdf_path)
    try:
        collection.delete(where={"source": filename})
    except Exception:
        pass

    pages = extract_document_data(pdf_path)
    chunks_list = []
    metas_list = []
    ids_list = []

    for page in pages:
        chunks = split_into_chunks(page["text"])
        for c_idx, chunk in enumerate(chunks):
            section_ref = detect_section_or_article(chunk, default_hint=page.get("section_hint", ""))
            chunks_list.append(chunk)
            metas_list.append({
                "source": filename,
                "page": page["page"],
                "section": section_ref,
                "chunk_id": c_idx
            })
            ids_list.append(f"doc_{filename}_p{page['page']}_c{c_idx}")

    if not chunks_list:
        return 0

    response = client_ai.embeddings.create(
        model="text-embedding-3-small",
        input=chunks_list
    )
    embeddings = [data.embedding for data in response.data]

    collection.add(
        ids=ids_list,
        embeddings=embeddings,
        documents=chunks_list,
        metadatas=metas_list
    )
    return len(chunks_list)

def restore_archived_documents(docs_dir: str = DOCS_FOLDER) -> List[str]:
    """Restaura los documentos eliminados desde la papelera local a DOCS."""
    restored = []
    if os.path.exists(ARCHIVE_FOLDER):
        for f in os.listdir(ARCHIVE_FOLDER):
            if is_supported_file(f):
                shutil.move(os.path.join(ARCHIVE_FOLDER, f), os.path.join(docs_dir, f))
                restored.append(f)
    return restored

def get_indexed_sources() -> List[str]:
    """Retorna lista de nombres de archivo ya indexados en ChromaDB."""
    try:
        chroma_client = chromadb.PersistentClient(path=CHROMA_DATA_PATH)
        col = chroma_client.get_collection(name="esan_reglamentos")
        result = col.get(include=["metadatas"])
        sources = list({m["source"] for m in result["metadatas"] if "source" in m})
        return sources
    except Exception:
        return []

def index_missing_documents(api_key: str, docs_dir: str = DOCS_FOLDER, progress_callback=None) -> Dict[str, Any]:
    """
    Indexación incremental inteligente:
    Detecta qué archivos de la carpeta local aún NO están en ChromaDB
    y SOLO procesa e indexa los nuevos/faltantes.
    NO re-indexa ni vuelve a gastar tokens en los documentos que ya existen.
    """
    all_paths = get_all_supported_files(docs_dir)
    indexed_set = set(get_indexed_sources())

    missing = [p for p in all_paths if os.path.basename(p) not in indexed_set]

    if not missing:
        return {
            "indexed_count": 0,
            "total_chunks": 0,
            "already_indexed_count": len(indexed_set),
            "new_files": [],
            "message": "Todos los documentos ya están indexados. No hay archivos nuevos pendientes."
        }

    total_chunks = 0
    new_files = []
    for idx, path in enumerate(missing, start=1):
        filename = os.path.basename(path)
        if progress_callback:
            progress_callback(f"Indexando archivo nuevo [{idx}/{len(missing)}]: {filename}")
        chunks = index_single_pdf(api_key, path)
        total_chunks += chunks
        new_files.append(filename)

    return {
        "indexed_count": len(missing),
        "total_chunks": total_chunks,
        "already_indexed_count": len(indexed_set),
        "new_files": new_files,
        "message": f"Se indexaron {len(missing)} documento(s) nuevo(s) ({total_chunks} secciones)."
    }

def build_vector_store(api_key: str, docs_dir: str = DOCS_FOLDER, progress_callback=None):
    """Procesa todos los archivos (.pdf, .docx, .xlsx) y crea los embeddings en ChromaDB."""
    client_ai = OpenAI(api_key=api_key)
    chroma_client = chromadb.PersistentClient(path=CHROMA_DATA_PATH)
    
    collection_name = "esan_reglamentos"
    try:
        chroma_client.delete_collection(name=collection_name)
    except Exception:
        pass
        
    collection = chroma_client.create_collection(
        name=collection_name,
        metadata={"hnsw:space": "cosine"}
    )
    
    os.makedirs(docs_dir, exist_ok=True)
    doc_paths = get_all_supported_files(docs_dir)
    
    total_files = len(doc_paths)
    all_chunks = []
    all_metadatas = []
    all_ids = []
    
    doc_counter = 0
    for idx, file_path in enumerate(doc_paths):
        doc_name = os.path.basename(file_path)
        if progress_callback:
            progress_callback(f"Extrayendo documento [{idx+1}/{total_files}]: {doc_name}")
            
        pages = extract_document_data(file_path)
        for page in pages:
            chunks = split_into_chunks(page["text"])
            for c_idx, chunk in enumerate(chunks):
                doc_counter += 1
                section_ref = detect_section_or_article(chunk, default_hint=page.get("section_hint", ""))
                all_chunks.append(chunk)
                all_metadatas.append({
                    "source": page["filename"],
                    "page": page["page"],
                    "section": section_ref,
                    "chunk_id": c_idx
                })
                all_ids.append(f"doc_{doc_counter}")

    if not all_chunks:
        raise ValueError("No se encontraron fragmentos de texto en los documentos.")

    batch_size = 80
    total_batches = (len(all_chunks) + batch_size - 1) // batch_size
    
    for i in range(0, len(all_chunks), batch_size):
        batch_chunks = all_chunks[i:i + batch_size]
        batch_meta = all_metadatas[i:i + batch_size]
        batch_ids = all_ids[i:i + batch_size]
        
        batch_num = (i // batch_size) + 1
        if progress_callback:
            progress_callback(f"Indexando vectores: lote {batch_num}/{total_batches} ({len(batch_chunks)} secciones)...")
            
        response = client_ai.embeddings.create(
            model="text-embedding-3-small",
            input=batch_chunks
        )
        embeddings = [data.embedding for data in response.data]
        
        collection.add(
            ids=batch_ids,
            embeddings=embeddings,
            documents=chunks,
            metadatas=batch_meta
        ) if False else collection.add(
            ids=batch_ids,
            embeddings=embeddings,
            documents=batch_chunks,
            metadatas=batch_meta
        )
        
    return {
        "total_files": total_files,
        "total_chunks": len(all_chunks),
        "collection_name": collection_name
    }

if __name__ == "__main__":
    from dotenv import load_dotenv
    load_dotenv()
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        print("ERROR: Define OPENAI_API_KEY en tu entorno o archivo .env")
    else:
        print(f"Indexando documentos desde {DOCS_FOLDER}...")
        res = build_vector_store(api_key, progress_callback=print)
        print(f"Listo! Se indexaron {res['total_files']} documentos ({res['total_chunks']} secciones).")
